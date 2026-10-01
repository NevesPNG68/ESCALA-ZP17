from __future__ import annotations

import json
import os
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook


BASE = Path(__file__).resolve().parents[1]
STATUS_FILE = BASE / "APP" / "data" / "update-status.json"
WORKBOOK = BASE / "PLANILHA" / "Escala_ZP17_2026.xlsx"
LOCAL_TZ = ZoneInfo("America/Sao_Paulo")
TARGET_YEAR = 2026


def last_day_of_month(current: date) -> date:
    if current.month == 12:
        following = date(current.year + 1, 1, 1)
    else:
        following = date(current.year, current.month + 1, 1)
    return following - timedelta(days=1)


def next_competence(current: date) -> str:
    if current.month == 12:
        return f"{current.year + 1}-01"
    return f"{current.year}-{current.month + 1:02d}"


def published_competences() -> set[str]:
    if not WORKBOOK.exists():
        return set()
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    try:
        sheet = workbook["PUBLICACOES"]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return set()
    headers = [str(value or "").strip() for value in rows[0]]
    positions = {header: index for index, header in enumerate(headers)}
    competence_index = positions.get("COMPETENCIA")
    status_index = positions.get("STATUS")
    if competence_index is None or status_index is None:
        return set()
    return {
        str(row[competence_index])
        for row in rows[1:]
        if len(row) > max(competence_index, status_index)
        and row[competence_index]
        and row[status_index] != "NAO_LOCALIZADO"
    }


def last_checked_date() -> date | None:
    try:
        value = json.loads(STATUS_FILE.read_text(encoding="utf-8")).get("checked_at")
        if not value:
            return None
        return datetime.fromisoformat(str(value).replace("Z", "+00:00")).astimezone(LOCAL_TZ).date()
    except (OSError, ValueError, TypeError, json.JSONDecodeError):
        return None


def is_due(current: date, checked_on: date | None, next_is_published: bool) -> tuple[bool, str]:
    if current.year != TARGET_YEAR:
        return False, "fora do escopo 2026"

    month_end = last_day_of_month(current)
    critical_start = month_end - timedelta(days=5)
    if current >= critical_start:
        return True, "janela crítica: verificação diária"

    elapsed = (current - checked_on).days if checked_on else None
    if current.day <= 10 and not next_is_published:
        if current.day in {1, 4, 7, 10} or elapsed is None or elapsed >= 3:
            return True, "escala do mês seguinte ausente: verificação a cada 3 dias"
        return False, "aguardando o intervalo de 3 dias"

    if current.day in {1, 11, 21} or elapsed is None or elapsed >= 10:
        return True, "fora da janela crítica: verificação a cada 10 dias"
    return False, "aguardando o intervalo de 10 dias"


def write_output(due: bool, reason: str) -> None:
    output = os.environ.get("GITHUB_OUTPUT")
    if output:
        with open(output, "a", encoding="utf-8") as stream:
            stream.write(f"due={'true' if due else 'false'}\n")
            stream.write(f"reason={reason}\n")


def main() -> int:
    today = datetime.now(LOCAL_TZ).date()
    competence = next_competence(today)
    available = competence in published_competences()
    due, reason = is_due(today, last_checked_date(), available)
    write_output(due, reason)
    print(f"AGENDAMENTO: {'EXECUTAR' if due else 'PULAR'} | {today.isoformat()} | {reason}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
